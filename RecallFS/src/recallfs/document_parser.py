"""文档解析：文本/代码/日志，以及可选的 PDF、DOCX、PPTX。

诚实降级原则
------------
* 可选依赖缺失 → ``status="unsupported"``，只保留元数据记录，并在错误字段说明原因；
* 扫描版 PDF 无文本层 → ``unsupported``，明确提示需要 OCR，绝不编造内容；
* 解析异常 → ``status="error"``，错误原因写入索引，可重试；
* 任何解析器都不写回原文件。
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .filesystem_watcher import read_limited

DEFAULT_MAX_BYTES = 10 * 1024 * 1024
DEFAULT_MAX_PDF_PAGES = 200

TEXT_EXTENSIONS: tuple[str, ...] = (
    ".txt",
    ".md",
    ".markdown",
    ".log",
    ".rst",
    ".csv",
    ".tsv",
    ".json",
    ".jsonl",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    ".conf",
    ".env",
    ".py",
    ".pyi",
    ".java",
    ".js",
    ".jsx",
    ".mjs",
    ".cjs",
    ".ts",
    ".tsx",
    ".c",
    ".h",
    ".cpp",
    ".cc",
    ".cxx",
    ".hpp",
    ".sql",
    ".sh",
    ".bash",
    ".zsh",
    ".ps1",
    ".bat",
    ".cmd",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".kt",
    ".swift",
    ".xml",
    ".html",
    ".htm",
    ".css",
    ".scss",
    ".less",
    ".tex",
    ".gitignore",
    ".dockerignore",
    ".editorconfig",
    ".properties",
    ".gradle",
    ".make",
)

IMAGE_EXTENSIONS: tuple[str, ...] = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp", ".heic")
OFFICE_EXTENSIONS: tuple[str, ...] = (".pdf", ".docx", ".pptx")
SUPPORTED_EXTENSIONS: tuple[str, ...] = TEXT_EXTENSIONS + OFFICE_EXTENSIONS


@dataclass
class ParseResult:
    status: str  # ok / unsupported / error / incomplete
    parser: str
    kind: str = "raw_text"  # raw_text / metadata
    pages: list[tuple[int | None, str]] = field(default_factory=list)
    text: str = ""
    error: str | None = None
    truncated: bool = False
    encoding: str | None = None
    notes: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.status == "ok"

    def to_dict(self) -> dict[str, Any]:
        return {
            "status": self.status,
            "parser": self.parser,
            "kind": self.kind,
            "text_length": len(self.text),
            "pages": len(self.pages),
            "error": self.error,
            "truncated": self.truncated,
            "encoding": self.encoding,
            "notes": list(self.notes),
        }


def pdf_available() -> bool:
    try:
        import pypdf  # noqa: F401
    except Exception:
        return False
    return True


def docx_available() -> bool:
    try:
        import docx  # noqa: F401
    except Exception:
        return False
    return True


def pptx_available() -> bool:
    try:
        import pptx  # noqa: F401
    except Exception:
        return False
    return True


def optional_parsers_status() -> dict[str, bool]:
    return {"pypdf": pdf_available(), "python-docx": docx_available(), "python-pptx": pptx_available()}


_BOMS: tuple[tuple[bytes, str], ...] = (
    (b"\xef\xbb\xbf", "utf-8-sig"),
    (b"\xff\xfe\x00\x00", "utf-32"),
    (b"\x00\x00\xfe\xff", "utf-32"),
    (b"\xff\xfe", "utf-16"),
    (b"\xfe\xff", "utf-16"),
)


def decode_text(data: bytes, *, truncated: bool = False) -> tuple[str, str]:
    """按 BOM → utf-8 → gb18030 → latin-1 的顺序解码，返回 ``(text, encoding)``。

    ``truncated=True`` 表示字节流可能在多字节字符中间被切断，此时对末尾残缺字节使用
    ``errors="ignore"``，避免整段文本误回退到 latin-1 而变成乱码。
    """
    for bom, encoding in _BOMS:
        if data.startswith(bom):
            try:
                return data.decode(encoding), encoding
            except UnicodeDecodeError:
                break
    for encoding in ("utf-8", "gb18030"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            if truncated:
                try:
                    return data.decode(encoding, errors="ignore"), encoding
                except UnicodeDecodeError:
                    continue
            continue
    return data.decode("latin-1", errors="replace"), "latin-1"


def _looks_binary(data: bytes) -> bool:
    if not data:
        return False
    sample = data[:8192]
    if b"\x00" in sample:
        return True
    printable = sum(1 for byte in sample if 32 <= byte < 127 or byte in (9, 10, 13) or byte >= 128)
    return printable / max(1, len(sample)) < 0.75


def _parse_plain_text(path: str, max_bytes: int) -> ParseResult:
    data, truncated = read_limited(path, max_bytes)
    if _looks_binary(data):
        return ParseResult(
            status="unsupported",
            parser="text",
            kind="metadata",
            error="文件疑似二进制（含 NUL 字节），未按文本解析",
        )
    text, encoding = decode_text(data, truncated=truncated)
    return ParseResult(
        status="ok",
        parser="text",
        pages=[(None, text)],
        text=text,
        truncated=truncated,
        encoding=encoding,
        notes=["仅索引文件前 10MB" ] if truncated else [],
    )


def _parse_pdf(path: str, max_pages: int) -> ParseResult:
    if not pdf_available():
        return ParseResult(
            status="unsupported",
            parser="pypdf",
            kind="metadata",
            error="pypdf 不可用（可选依赖未安装），PDF 仅保留元数据记录",
        )
    import pypdf

    try:
        reader = pypdf.PdfReader(path)
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception as exc:
                return ParseResult(
                    status="error",
                    parser="pypdf",
                    kind="metadata",
                    error=f"PDF 已加密且无法解密: {exc}",
                )
        total_pages = len(reader.pages)
        limit = min(total_pages, max_pages)
        pages: list[tuple[int | None, str]] = []
        notes: list[str] = []
        for index in range(limit):
            try:
                page_text = reader.pages[index].extract_text() or ""
            except Exception as exc:
                notes.append(f"第 {index + 1} 页文本提取失败: {exc}")
                page_text = ""
            pages.append((index + 1, page_text))
        text = "\n".join(page_text for _no, page_text in pages)
        truncated = total_pages > max_pages
        if not text.strip():
            return ParseResult(
                status="unsupported",
                parser="pypdf",
                kind="metadata",
                pages=pages,
                text="",
                truncated=truncated,
                error="PDF 无可提取文本层（可能为扫描件/图片），需要 OCR",
                notes=notes,
            )
        return ParseResult(
            status="ok",
            parser="pypdf",
            kind="raw_text",
            pages=pages,
            text=text,
            truncated=truncated,
            notes=notes + ([f"仅解析前 {max_pages} 页（共 {total_pages} 页）"] if truncated else []),
        )
    except Exception as exc:
        return ParseResult(
            status="error",
            parser="pypdf",
            kind="metadata",
            error=f"PDF 解析失败: {type(exc).__name__}: {exc}",
        )


def _parse_docx(path: str, max_bytes: int) -> ParseResult:
    if not docx_available():
        return ParseResult(
            status="unsupported",
            parser="python-docx",
            kind="metadata",
            error="python-docx 不可用（可选依赖未安装），DOCX 仅保留元数据记录",
        )
    import docx

    try:
        document = docx.Document(path)
        parts: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        truncated = len(text.encode("utf-8", errors="ignore")) > max_bytes
        if truncated:
            text = text[: max_bytes // 2]
        return ParseResult(
            status="ok",
            parser="python-docx",
            pages=[(None, text)],
            text=text,
            truncated=truncated,
            encoding="utf-8",
            notes=["DOCX 无固定页码，按段落行号定位"] + (["内容超限已截断"] if truncated else []),
        )
    except Exception as exc:
        return ParseResult(
            status="error",
            parser="python-docx",
            kind="metadata",
            error=f"DOCX 解析失败: {type(exc).__name__}: {exc}",
        )


def _parse_pptx(path: str) -> ParseResult:
    if not pptx_available():
        return ParseResult(
            status="unsupported",
            parser="python-pptx",
            kind="metadata",
            error="python-pptx 不可用（可选依赖未安装），PPTX 仅保留元数据记录",
        )
    import pptx

    try:
        presentation = pptx.Presentation(path)
        pages: list[tuple[int | None, str]] = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    try:
                        if shape.text_frame.text.strip():
                            texts.append(shape.text_frame.text)
                    except Exception:
                        continue
                if getattr(shape, "has_table", False):
                    try:
                        for row in shape.table.rows:
                            texts.append(" | ".join(cell.text for cell in row.cells))
                    except Exception:
                        continue
            pages.append((index, "\n".join(texts)))
        text = "\n".join(page_text for _no, page_text in pages)
        if not text.strip():
            return ParseResult(
                status="unsupported",
                parser="python-pptx",
                kind="metadata",
                pages=pages,
                error="PPTX 未提取到文本（可能全部为图片），需要 OCR",
            )
        return ParseResult(status="ok", parser="python-pptx", pages=pages, text=text, encoding="utf-8")
    except Exception as exc:
        return ParseResult(
            status="error",
            parser="python-pptx",
            kind="metadata",
            error=f"PPTX 解析失败: {type(exc).__name__}: {exc}",
        )


def parse_document(
    path: str | os.PathLike[str],
    *,
    max_bytes: int = DEFAULT_MAX_BYTES,
    max_pdf_pages: int = DEFAULT_MAX_PDF_PAGES,
) -> ParseResult:
    """按扩展名解析文件。永远返回 ParseResult，不抛出文件内容相关异常。"""
    fs_path = os.fspath(path)
    ext = Path(fs_path).suffix.casefold()
    try:
        size = os.path.getsize(fs_path)
    except OSError as exc:
        return ParseResult(status="error", parser="none", kind="metadata", error=f"无法读取文件: {exc}")

    if ext in IMAGE_EXTENSIONS:
        return ParseResult(
            status="unsupported",
            parser="none",
            kind="metadata",
            error="图片文件不解析，需要 OCR（本版本不内置 OCR）",
        )
    if ext == ".pdf":
        return _parse_pdf(fs_path, max_pdf_pages)
    if ext == ".docx":
        return _parse_docx(fs_path, max_bytes)
    if ext == ".pptx":
        return _parse_pptx(fs_path)
    if ext in TEXT_EXTENSIONS:
        return _parse_plain_text(fs_path, max_bytes)
    if ext == "":
        data, _truncated = read_limited(fs_path, min(max_bytes, 64 * 1024))
        if not _looks_binary(data):
            return _parse_plain_text(fs_path, max_bytes)
    return ParseResult(
        status="unsupported",
        parser="none",
        kind="metadata",
        error=f"没有可用解析器: 扩展名 {ext or '(无)'} 不在支持列表内" + (f"，文件大小 {size} 字节" if size else ""),
    )


def live_text_for_evidence(
    path: str | os.PathLike[str],
    *,
    page: int | None = None,
    max_bytes: int = 16 * 1024 * 1024,
) -> tuple[str, str | None]:
    """读取“当前磁盘内容”用于证据校验。返回 ``(text, error)``。"""
    fs_path = os.fspath(path)
    ext = Path(fs_path).suffix.casefold()
    if ext == ".pdf":
        if not pdf_available():
            return "", "pypdf 不可用，无法校验 PDF 当前内容"
        import pypdf

        try:
            reader = pypdf.PdfReader(fs_path)
            if page is not None and 0 < page <= len(reader.pages):
                return reader.pages[page - 1].extract_text() or "", None
            return "\n".join((p.extract_text() or "") for p in reader.pages[:200]), None
        except Exception as exc:
            return "", f"PDF 重读失败: {exc}"
    try:
        data, truncated = read_limited(fs_path, max_bytes)
        text, _encoding = decode_text(data, truncated=truncated)
        return text, None
    except OSError as exc:
        return "", f"重读失败: {exc}"


_REFERENCE_RE = re.compile(r"[\w./\\-]+\.(?:py|js|ts|tsx|jsx|java|c|cpp|h|sql|md|txt|json|yaml|yml|toml|ini|sh|pdf|docx|pptx)\b", re.I)


def extract_referenced_paths(text: str) -> list[str]:
    """从文本中抽取形如路径/文件名的引用（用于相关文件关联）。"""
    return list(dict.fromkeys(match.group(0) for match in _REFERENCE_RE.finditer(text)))
