from __future__ import annotations

import pytest

from helpers import make_pdf, write_file

import recallfs.document_parser as document_parser
import recallfs.index_scheduler as index_scheduler
from recallfs.document_parser import ParseResult, optional_parsers_status, parse_document, pdf_available
from recallfs.index_scheduler import IndexScheduler
from recallfs.search_engine import SearchFilters

PLAIN_EXTENSIONS = [".md", ".py", ".json", ".yaml", ".yml", ".toml", ".ini", ".log", ".sql", ".sh", ".c", ".cpp", ".h", ".js", ".ts", ".java"]


def test_all_plain_text_formats(corpus, scheduler, db):
    for index, ext in enumerate(PLAIN_EXTENSIONS):
        write_file(corpus, f"sample_{index}{ext}", f"通用内容 format_marker_{index}\n")
    stats = scheduler.scan([str(corpus)])
    assert stats.added == len(PLAIN_EXTENSIONS)
    assert stats.parse_errors == 0
    rows = db.all_files()
    assert all(row["parse_status"] == "ok" for row in rows)
    assert all(row["parser"] == "text" for row in rows)


def test_pdf_pages_and_evidence(corpus, scheduler, engine, db):
    if not pdf_available():
        pytest.skip("pypdf 不可用")
    path = make_pdf(
        corpus / "manual.pdf",
        ["RecallFS page one alpha", "Second page contains evidencekeyword"],
    )
    scheduler.scan([str(corpus)])
    row = db.get_file_by_path(path)
    assert row["parse_status"] == "ok"
    assert row["parser"] == "pypdf"
    pages = {chunk["page"] for chunk in db.chunks_for_file(int(row["id"]))}
    assert {1, 2}.issubset(pages)

    result = engine.search("evidencekeyword")
    assert result.hits
    hit = result.hits[0]
    assert hit.page == 2
    assert hit.path == str(path)
    evidence = engine.evidence(hit)
    assert evidence["stale"] is False
    assert evidence["page"] == 2
    assert evidence["content_verified"] is True


def test_broken_pdf_is_error_and_visible(corpus, scheduler, engine, db):
    path = write_file(corpus, "broken_manual.pdf", b"%PDF-1.4\nnot really a pdf\n")
    scheduler.scan([str(corpus)])
    row = db.get_file_by_path(path)
    assert row is not None
    assert row["parse_status"] == "error"
    assert row["parse_error"]
    assert "PDF" in row["parse_error"]
    result = engine.search("broken_manual")
    assert result.hits
    assert result.hits[0].layer == "exact_name"
    assert result.hits[0].parse_status == "error"
    assert result.hits[0].parse_error


def test_pdf_without_pypdf_is_unsupported(corpus, db, monkeypatch):
    monkeypatch.setattr(document_parser, "pdf_available", lambda: False)
    path = make_pdf(corpus / "needs_pypdf.pdf", ["text here"])
    scheduler = IndexScheduler(db)
    stats = scheduler.scan([str(corpus)])
    row = db.get_file_by_path(path)
    assert row is not None
    assert row["parse_status"] == "unsupported"
    assert "pypdf" in row["parse_error"]
    assert stats.unsupported == 1
    direct = parse_document(path)
    assert direct.status == "unsupported"


def test_unsupported_extension_keeps_record(corpus, scheduler, engine, db):
    path = write_file(corpus, "archive.xyz", "FakeContent 不应被当作文本读取\n")
    scheduler.scan([str(corpus)])
    row = db.get_file_by_path(path)
    assert row is not None
    assert row["parse_status"] == "unsupported"
    assert "没有可用解析器" in row["parse_error"]
    assert engine.search("FakeContent").empty
    named = engine.search("archive.xyz")
    assert named.hits
    assert named.hits[0].parse_status == "unsupported"


def test_image_requires_ocr(corpus, scheduler, engine, db):
    path = write_file(corpus, "scan.png", b"\x89PNG\r\n\x1a\nnot really an image")
    scheduler.scan([str(corpus)])
    row = db.get_file_by_path(path)
    assert row["parse_status"] == "unsupported"
    assert "OCR" in row["parse_error"]
    assert engine.search("not really").empty


def test_truncation_marks_file_and_limits_content(corpus, db):
    content = "开头标记 " + "填充内容 " * 500 + "结尾标记"
    path = write_file(corpus, "large.txt", content)
    scheduler = IndexScheduler(db, max_bytes=200)
    scheduler.scan([str(corpus)])
    row = db.get_file_by_path(path)
    assert row["truncated"] == 1
    assert row["parse_status"] == "ok"
    from recallfs.search_engine import SearchEngine

    search = SearchEngine(db)
    assert search.search("开头标记").hits
    assert search.search("结尾标记").empty


def test_docx_pptx_available_or_marked(corpus, db, monkeypatch):
    import recallfs.document_parser as parser

    if parser.docx_available():
        import docx

        document = docx.Document()
        document.add_paragraph("DOCX 独特关键词 段落内容")
        document.add_paragraph("第二段")
        docx_path = corpus / "manual.docx"
        document.save(str(docx_path))
        IndexScheduler(db).scan([str(corpus)])
        row = db.get_file_by_path(docx_path)
        assert row["parse_status"] == "ok"
        from recallfs.search_engine import SearchEngine

        assert SearchEngine(db).search("独特关键词").hits
    else:
        docx_path = write_file(corpus, "manual.docx", b"not a docx")
        monkeypatch.setattr(parser, "docx_available", lambda: False)
        IndexScheduler(db).scan([str(corpus)])
        row = db.get_file_by_path(docx_path)
        assert row["parse_status"] == "unsupported"

    if parser.pptx_available():
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        textbox.text_frame.text = "PPTX 幻灯片关键词"
        pptx_path = corpus / "slides.pptx"
        presentation.save(str(pptx_path))
        IndexScheduler(db).scan([str(corpus)])
        row = db.get_file_by_path(pptx_path)
        assert row["parse_status"] == "ok"
        from recallfs.search_engine import SearchEngine

        hits = SearchEngine(db).search("幻灯片关键词").hits
        assert hits
        assert hits[0].page == 1


def test_incomplete_when_file_changes_during_index(corpus, db, monkeypatch):
    path = write_file(corpus, "changing.txt", "初始内容\n")

    calls = {"count": 0}

    def fake_parse(target, *, max_bytes=0, max_pdf_pages=0):
        calls["count"] += 1
        with open(target, "a", encoding="utf-8") as handle:
            handle.write(f"追加变更 {calls['count']}\n")
        return ParseResult(status="ok", parser="fake", pages=[(None, "内容")], text="内容")

    monkeypatch.setattr(index_scheduler, "parse_document", fake_parse)
    stats = IndexScheduler(db).scan([str(corpus)])
    row = db.get_file_by_path(path)
    assert row["parse_status"] == "incomplete"
    assert stats.incomplete == 1
    assert "incomplete" in (row["parse_error"] or "")
    assert calls["count"] >= 2


def test_optional_parser_status_flags():
    status = optional_parsers_status()
    assert set(status) == {"pypdf", "python-docx", "python-pptx"}
